#!/usr/bin/env python3
# Handling: Unclassified — working
"""Validate the Workforce Training and Simulation CSVs, then build the SQLite
database, the Excel workbook and the Seam coverage page.

    python tools/build.py            # validate and build
    python tools/build.py --check    # validate only

The CSV files in data/ are the source of truth. Everything in release/ is
generated: never edit it by hand, rebuild it.

Optional: set PRIVATE_TERMS to a comma-separated list of terms that must never
appear in the data (kept out of the repository on purpose).
"""
import csv, os, re, sqlite3, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA, DIST, SCHEMA = ROOT / "data", ROOT / "release", ROOT / "schema" / "schema.sql"
DB, XLSX, COVERAGE = DIST / "wts.sqlite", DIST / "WTS_Database.xlsx", DIST / "SEAM_COVERAGE.md"

# table -> (csv file, primary key, has source_id, has is_synthetic, has secondary cells)
TABLES = {
    "seam_cells":           ("seam_cells.csv", "cell_id", False, False, False),
    "sources":              ("sources.csv", "source_id", False, False, False),
    "programs":             ("programs.csv", "program_id", True, True, True),
    "facilities":           ("facilities.csv", "facility_id", True, True, True),
    "equipment":            ("equipment.csv", "equipment_id", False, True, False),
    "skills":               ("skills.csv", "skill_id", False, False, True),
    "skill_evidence":       ("skill_evidence.csv", "evidence_id", True, False, False),
    "roles":                ("roles.csv", "role_id", True, False, False),
    "indicators":           ("indicators.csv", "indicator_id", True, True, False),
    "enabling_environment": ("enabling_environment.csv", "instrument_id", True, False, True),
    "missions":             ("missions.csv", "mission_id", True, False, True),
    "role_skills":          ("role_skills.csv", "link_id", False, False, False),
}
INTEGER_COLS = {"layer_index", "capability_index", "is_synthetic", "quantity"}
REAL_COLS = {"value", "interval_months_low", "interval_months_high"}
DATE_RE = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")


def load():
    out = {}
    for t, (fn, *_rest) in TABLES.items():
        with open(DATA / fn, newline="", encoding="utf-8") as f:
            out[t] = list(csv.DictReader(f))
    return out


def validate(d):
    errs = []
    cells = {r["cell_id"] for r in d["seam_cells"]}
    sources = {r["source_id"] for r in d["sources"]}
    if len(d["seam_cells"]) != 60 or len(cells) != 60:
        errs.append("seam_cells must hold exactly 60 unique cells")
    for t, (fn, pk, has_src, has_syn, has_sec) in TABLES.items():
        seen = set()
        for i, r in enumerate(d[t], start=2):
            where = f"{fn} line {i}"
            rid = r[pk]
            if not rid:
                errs.append(f"{where}: empty {pk}")
            if rid in seen:
                errs.append(f"{where}: duplicate {pk} {rid}")
            seen.add(rid)
            if "cell_id" in r and t != "seam_cells" and r["cell_id"] not in cells:
                errs.append(f"{where}: unknown cell_id {r['cell_id']}")
            if has_sec:
                for c in filter(None, (r.get("secondary_cell_ids") or "").split(";")):
                    if c.strip() not in cells:
                        errs.append(f"{where}: unknown secondary cell {c}")
            for k in ("as_of", "accessed", "published", "issued"):
                if r.get(k) and not DATE_RE.match(r[k]):
                    errs.append(f"{where}: {k} '{r[k]}' is not YYYY, YYYY-MM or YYYY-MM-DD")
            syn = has_syn and r.get("is_synthetic") == "1"
            if has_syn and r.get("is_synthetic") not in ("0", "1"):
                errs.append(f"{where}: is_synthetic must be 0 or 1")
            # Synthetic rows must be unmistakable and must never carry a source.
            if syn:
                if not rid.startswith("SYN-"):
                    errs.append(f"{where}: synthetic row id must start with SYN-")
                name = r.get("name", "")
                if name and not name.startswith("Synthetic"):
                    errs.append(f"{where}: synthetic row name must start with 'Synthetic'")
                if r.get("source_id"):
                    errs.append(f"{where}: synthetic row must not cite a source")
            elif has_syn and rid.startswith("SYN-"):
                errs.append(f"{where}: SYN- id on a row not marked synthetic")
            # Public claims must be cited to a known source.
            if has_src and not syn:
                if not r.get("source_id"):
                    errs.append(f"{where}: public row has no source_id")
                elif r["source_id"] not in sources:
                    errs.append(f"{where}: unknown source_id {r['source_id']}")
    for r in d["sources"]:
        if not r["url"].startswith("https://"):
            errs.append(f"sources.csv {r['source_id']}: url must be https")
    fac = {r["facility_id"]: r for r in d["facilities"]}
    for r in d["equipment"]:
        if r["facility_id"] not in fac:
            errs.append(f"equipment.csv {r['equipment_id']}: unknown facility {r['facility_id']}")
        elif fac[r["facility_id"]]["is_synthetic"] != r["is_synthetic"]:
            errs.append(f"equipment.csv {r['equipment_id']}: synthetic flag differs from its facility")
    skills = {r["skill_id"] for r in d["skills"]}
    for r in d["skill_evidence"]:
        if r["skill_id"] not in skills:
            errs.append(f"skill_evidence.csv {r['evidence_id']}: unknown skill {r['skill_id']}")
    roles = {r["role_id"] for r in d["roles"]}
    for r in d["indicators"]:
        if r["role_id"] and r["role_id"] not in roles:
            errs.append(f"indicators.csv {r['indicator_id']}: unknown role {r['role_id']}")
    missions = {r["mission_id"] for r in d["missions"]}
    for r in d["role_skills"]:
        where = f"role_skills.csv {r['link_id']}"
        for col, known in (("mission_id", missions), ("role_id", roles), ("skill_id", skills)):
            if r[col] not in known:
                errs.append(f"{where}: unknown {col} {r[col]}")
        if r["basis"] == "source":
            if r["source_id"] not in sources:
                errs.append(f"{where}: basis 'source' needs a known source_id")
        elif r["basis"] == "inference":
            if r["source_id"]:
                errs.append(f"{where}: an inference must not cite a source; cite it and set basis to 'source'")
        else:
            errs.append(f"{where}: basis must be 'source' or 'inference'")
    terms = [t.strip().lower() for t in os.environ.get("PRIVATE_TERMS", "").split(",") if t.strip()]
    if terms:
        for fn in sorted(DATA.glob("*.csv")):
            text = fn.read_text(encoding="utf-8").lower()
            for t in terms:
                if t in text:
                    errs.append(f"{fn.name}: contains a private term (not printed)")
    return errs


def coerce(col, v):
    if v == "":
        return None
    if col in INTEGER_COLS:
        return int(v)
    if col in REAL_COLS:
        return float(v)
    return v


def build_sqlite(d):
    if DB.exists():
        DB.unlink()
    con = sqlite3.connect(DB)
    con.executescript(SCHEMA.read_text(encoding="utf-8"))
    for t in TABLES:
        rows = d[t]
        cols = list(rows[0].keys())
        con.executemany(
            f"INSERT INTO {t} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
            [[coerce(c, r[c]) for c in cols] for r in rows])
    links = []
    for t, (_fn, pk, _src, has_syn, has_sec) in TABLES.items():
        # skill_evidence and role_skills inherit their parents' bindings, so they add no cell links.
        if t in ("seam_cells", "sources", "skill_evidence", "role_skills"):
            continue
        for r in d[t]:
            syn = int(r.get("is_synthetic") or 0)
            links.append((t, r[pk], r["cell_id"], "primary", syn))
            if has_sec:
                for c in filter(None, (r.get("secondary_cell_ids") or "").split(";")):
                    links.append((t, r[pk], c.strip(), "secondary", syn))
    con.executemany("INSERT INTO record_cells VALUES (?,?,?,?,?)", links)
    con.executemany("INSERT INTO dataset_metadata VALUES (?,?)", [
        ("handling", "Unclassified — working"),
        ("licence_data", "OGL-Canada-2.0"), ("licence_code", "Apache-2.0"),
        ("source_of_truth", "data/*.csv"),
        ("substrate", "Sovereign Seam: 6 layers x 10 capabilities = 60 cells")])
    con.commit()
    bad = con.execute("PRAGMA foreign_key_check").fetchall()
    if bad:
        raise SystemExit(f"foreign key check failed: {bad[:5]}")
    cov = con.execute("SELECT * FROM v_seam_coverage").fetchall()
    con.execute("VACUUM")
    con.close()
    return cov


def coverage_grid(cov):
    caps = sorted({(r[0].split("-")[1], r[2]) for r in cov}, key=lambda x: int(x[0][1:]))
    layers = []
    for r in cov:
        if r[1] not in layers:
            layers.append(r[1])
    idx = {r[0]: r for r in cov}
    codes = {"Applications": "APP", "Sovereign Decision Layer": "SOV", "Models": "MOD",
             "Infrastructure": "INF", "Chips": "CHI", "Energy": "ENE"}
    return caps, layers, idx, codes


def write_coverage_md(cov):
    caps, layers, idx, codes = coverage_grid(cov)
    lines = ["Handling: Unclassified — working", "",
             "# Sovereign Seam coverage: Workforce Training and Simulation", "",
             "Generated by `tools/build.py`. Do not edit.", "",
             "Each cell shows **public records (primary + secondary bindings)**, then "
             "synthetic rows in brackets. Synthetic rows are for testing only and are "
             "never added to public counts. An empty cell is a gap in coverage, not a "
             "finding that nothing is there.", "",
             "| Layer | " + " | ".join(f"{c} {n}" for c, n in caps) + " |",
             "|---|" + "---|" * len(caps)]
    for L in layers:
        row = []
        for c, _n in caps:
            r = idx[f"{codes[L]}-{c}"]
            pub, syn = r[3] + r[4], r[5]
            cell = (f"**{pub}**" if pub else "") + (f" [{syn}]" if syn else "")
            row.append(cell.strip() or "·")
        lines.append(f"| {L} | " + " | ".join(row) + " |")
    pub_cells = sum(1 for r in cov if r[3] + r[4])
    lines += ["", f"Cells with at least one public record: **{pub_cells} of 60**.", ""]
    COVERAGE.write_text("\n".join(lines), encoding="utf-8")


def write_xlsx(d, cov):
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    teal, gold, grey = "137A88", "F0B429", "9AA7B4"
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=teal)
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme = [
        ["Workforce Training and Simulation database"],
        ["Element Ventures (Canada) Ltd. · Sovereign Operating Intelligence (SOI™)"],
        [""],
        ["What this is", "Public, cited reference data on training programs, simulation facilities, skills and retention evidence, workforce roles and indicators, and the enabling policy environment."],
        ["How it is organised", "Every record binds to one Sovereign Seam cell (6 stack layers x 10 Defence Industrial Strategy capabilities = 60 cells). See the 'Seam coverage' sheet."],
        ["Synthetic rows", "Rows with ids starting SYN- and names starting 'Synthetic' are invented for testing only. They are shaded grey and never counted with public data."],
        ["Citability", "P primary · A authoritative secondary · AS attributable secondary · AN attribute-only · PR self-reported · S single-sourced"],
        ["Source of truth", "The CSV files in data/. This workbook is generated by tools/build.py; edits made here are lost on the next build."],
        ["Licence", "Data: Open Government Licence - Canada 2.0. Code: Apache-2.0."],
        ["Handling", "Unclassified — working"],
    ]
    for r in readme:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14, color=teal)
    for row in ws.iter_rows(min_row=4, max_col=1):
        row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    cs = wb.create_sheet("Seam coverage")
    caps, layers, idx, codes = coverage_grid(cov)
    cs.append(["Layer"] + [f"{c} {n}" for c, n in caps])
    for L in layers:
        row = [L]
        for c, _n in caps:
            r = idx[f"{codes[L]}-{c}"]
            pub, syn = r[3] + r[4], r[5]
            row.append(pub if pub else (f"[{syn}]" if syn else None))
        cs.append(row)
    cs.append([])
    cs.append(["Public records per cell (primary + secondary). [n] = synthetic rows only. Blank = no coverage yet."])
    for c in cs[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cs.row_dimensions[1].height = 42
    cs.column_dimensions["A"].width = 26
    for i in range(2, 12):
        cs.column_dimensions[get_column_letter(i)].width = 13
    for row in cs.iter_rows(min_row=2, max_row=1 + len(layers), min_col=2):
        for c in row:
            c.alignment = Alignment(horizontal="center")
            if isinstance(c.value, int):
                c.fill = PatternFill("solid", fgColor=gold)
                c.font = Font(bold=True)

    syn_fill = PatternFill("solid", fgColor="E6EAEE")
    for t in TABLES:
        sh = wb.create_sheet(t[:31])
        rows = d[t]
        cols = list(rows[0].keys())
        sh.append(cols)
        for r in rows:
            sh.append([coerce(c, r[c]) for c in cols])
            if r.get("is_synthetic") == "1":
                for c in sh[sh.max_row]:
                    c.fill = syn_fill
        for c in sh[1]:
            c.font, c.fill = head_font, head_fill
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for i, col in enumerate(cols, start=1):
            width = max(len(col), *(len(str(r[col])) for r in rows))
            sh.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)

    for sh in wb.worksheets:  # print landscape, one page wide
        sh.page_setup.orientation = "landscape"
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0

    fixed = datetime(2026, 1, 1)  # fixed metadata keeps rebuilds byte-stable
    wb.properties.creator = "Element Ventures (Canada) Ltd."
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.save(XLSX)
    # openpyxl stamps each zip entry with the save time; rewrite with a fixed
    # stamp so an unchanged dataset rebuilds to an identical file.
    import zipfile
    with zipfile.ZipFile(XLSX) as z:
        entries = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(XLSX, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in entries:
            if info.filename == "docProps/core.xml":  # openpyxl resets 'modified' on save
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type, zi.external_attr = zipfile.ZIP_DEFLATED, info.external_attr
            z.writestr(zi, data)


def main():
    d = load()
    errs = validate(d)
    if errs:
        print(f"FAILED: {len(errs)} problem(s)")
        for e in errs:
            print("  -", e)
        sys.exit(1)
    counts = {t: len(d[t]) for t in TABLES}
    syn = sum(1 for t in TABLES for r in d[t] if r.get("is_synthetic") == "1")
    print(f"Validation passed: {sum(counts.values())} rows ({syn} synthetic) across {len(TABLES)} tables.")
    if "--check" in sys.argv:
        return
    DIST.mkdir(exist_ok=True)
    cov = build_sqlite(d)
    write_coverage_md(cov)
    write_xlsx(d, cov)
    for t, n in counts.items():
        print(f"  {t:<22} {n}")
    print(f"Built {DB.relative_to(ROOT)}, {XLSX.relative_to(ROOT)}, {COVERAGE.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
